"""Exports (§9 r37, §17.11): the Registrations sheet in Paperform's own
column order with our check-in and redeemed columns filled, bookings, claims,
and a printable fallback list for when the laptop is down."""

import html
import io
from datetime import datetime

import openpyxl

import config
from services import bookings, claims, roster

KINDS = ("registrations", "bookings", "claims", "fallback")
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _when(utc_iso):
    if not utc_iso:
        return ""
    return datetime.fromisoformat(utc_iso).astimezone(config.TIMEZONE).strftime("%Y-%m-%d %H:%M")


def _handle(r):
    return "@" + r["handle"] if r["handle"] else ""


def _book(sheets):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, header, rows in sheets:
        ws = wb.create_sheet(title)
        ws.append(header)
        for row in rows:
            ws.append(list(row))
        for i, name in enumerate(header, start=1):
            width = max([len(str(name))] + [len(str(r[i - 1] or "")) for r in rows] + [8])
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = min(width + 2, 60)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _people(conn):
    return conn.execute("SELECT * FROM attendees WHERE is_test=0 ORDER BY id").fetchall()


def _claims_by_person(conn):
    out = {}
    for c in conn.execute("SELECT * FROM claims WHERE voided_at IS NULL ORDER BY claimed_at"):
        out.setdefault(c["attendee_id"], []).append(c)
    return out


def _item(c):
    """"Pastry (Mini Tart)": which kind went, when the item has kinds."""
    kind = claims.variant_label(c["item"], c["variant"]) if c["variant"] else ""
    return claims.label(c["item"]) + (f" ({kind})" if kind else "")


def registrations(conn):
    got = _claims_by_person(conn)
    rows = []
    for p in _people(conn):
        mine = [c for c in got.get(p["id"], []) if c["item"] in claims.ITEMS]
        rows.append((
            p["submitted_at_text"] or "", p["paperform_id"] or "", p["name"],
            p["handle_raw"] or _handle(p), p["email"] or "", p["paperform_receipt_url"] or "",
            "Yes" if p["checked_in_at"] else "", _when(p["checked_in_at"]),
            ", ".join(_item(c) for c in mine), _when(mine[-1]["claimed_at"]) if mine else "",
        ))
    return _book([("Registrations", roster.COLUMNS, rows)])


def bookings_sheet(conn):
    escape = conn.execute(
        "SELECT b.*, s.starts_at, a.name, a.handle, o.handle AS owner FROM escape_bookings b "
        "JOIN slots s ON s.id=b.slot_id JOIN attendees a ON a.id=b.attendee_id "
        "LEFT JOIN attendees o ON o.id=b.booked_by_id ORDER BY s.starts_at, b.id").fetchall()
    # The jam room is a group booking too since 18 Sep, so the sheet carries
    # who booked it — without that a slot reads as several unrelated people.
    jam = conn.execute(
        "SELECT j.*, s.starts_at, s.ends_at, a.name, a.handle, o.handle AS owner "
        "FROM jam_bookings j JOIN slots s ON s.id=j.slot_id "
        "JOIN attendees a ON a.id=j.attendee_id "
        "LEFT JOIN attendees o ON o.id=j.booked_by_id ORDER BY s.starts_at, j.id").fetchall()
    return _book([
        ("Escape room", ["Ref", "Game", "Name", "Username", "Status", "Booked by", "Booked at"],
         [(b["ref_code"], claims.clock(b["starts_at"]), b["name"], _handle(b),
           b["status"], "@" + b["owner"] if b["owner"] else "", _when(b["created_at"])) for b in escape]),
        ("Jamming studio",
         ["Ref", "Slot", "Name", "Username", "Status", "Booked by", "Booked at"],
         [(j["ref_code"], f"{claims.clock(j['starts_at'])}–{claims.clock(j['ends_at'])}", j["name"],
           _handle(j), j["status"], "@" + j["owner"] if j["owner"] else "",
           _when(j["created_at"])) for j in jam]),
    ])


def claims_sheet(conn):
    rows = conn.execute(
        "SELECT c.*, a.name, a.handle FROM claims c JOIN attendees a ON a.id=c.attendee_id "
        "ORDER BY c.claimed_at").fetchall()
    checked = conn.execute("SELECT * FROM attendees WHERE checked_in_at IS NOT NULL "
                           "ORDER BY checked_in_at").fetchall()
    return _book([
        ("Claims", ["Time", "Name", "Username", "Item", "Staff", "Station", "Voided at", "Void reason"],
         [(_when(c["claimed_at"]), c["name"], _handle(c), _item(c), c["staff_name"] or "",
           c["station"] or "", _when(c["voided_at"]), c["void_reason"] or "") for c in rows]),
        ("Check-ins", ["Time", "Name", "Username", "Checked in by"],
         [(_when(p["checked_in_at"]), p["name"], _handle(p), p["checked_in_by"] or "") for p in checked]),
    ])


def fallback(conn):
    """One printable page: everyone, their pass code, and what they have."""
    got = _claims_by_person(conn)
    lines = []
    for p in conn.execute("SELECT * FROM attendees WHERE status='active' ORDER BY name COLLATE NOCASE"):
        b = bookings.active_booking(conn, p["id"])
        jams = bookings.jam_bookings_of(conn, p["id"])
        have = {c["item"] for c in got.get(p["id"], [])}
        cells = [p["name"], _handle(p), p["pass_code"] or "", p["payment_status"],
                 " ".join(("■ " if k in have else "□ ") + v for k, v in config.ITEMS),
                 claims.clock(b["starts_at"]) if b else "",
                 ", ".join(claims.clock(j["starts_at"]) for j in jams)]
        lines.append("<tr>" + "".join(f"<td>{html.escape(str(x))}</td>" for x in cells) + "</tr>")
    head = "".join(f"<th>{h}</th>" for h in
                   ("Name", "Username", "Pass", "Payment", "Collected", "Escape", "Jam"))
    stamp = datetime.now(config.TIMEZONE).strftime("%d %b %Y, %H:%M")
    page = f"""<!DOCTYPE html><html lang="en"><head><meta charset="utf-8">
<title>The Yard — fallback list</title>
<style>
body{{font:12px/1.35 Arial,sans-serif;margin:18px;color:#111}}
h1{{font-size:18px;margin:0 0 4px}} p{{margin:0 0 12px;color:#444}}
table{{border-collapse:collapse;width:100%}} th,td{{border:1px solid #999;padding:4px 6px;text-align:left;vertical-align:top}}
th{{background:#eee}} tr{{page-break-inside:avoid}}
@media print{{body{{margin:0}}}}
</style></head><body>
<h1>The Yard — offline fallback list</h1>
<p>Printed {html.escape(stamp)}. Tick items by hand while the laptop is down, and enter them afterwards.</p>
<table><thead><tr>{head}</tr></thead><tbody>{''.join(lines)}</tbody></table>
<script>window.print()</script>
</body></html>"""
    return page.encode("utf-8")


def build(conn, kind):
    """(bytes, file name, mime type)."""
    day = config.EVENT_DATE
    if kind == "registrations":
        return registrations(conn), f"The_Yard_Registrations_{day}.xlsx", XLSX
    if kind == "bookings":
        return bookings_sheet(conn), f"The_Yard_Bookings_{day}.xlsx", XLSX
    if kind == "claims":
        return claims_sheet(conn), f"The_Yard_Claims_{day}.xlsx", XLSX
    if kind == "fallback":
        return fallback(conn), "The_Yard_fallback.html", "text/html; charset=utf-8"
    raise ValueError(kind)
