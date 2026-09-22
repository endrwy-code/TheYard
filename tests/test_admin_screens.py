"""Overview, Settings, Audit, backups, exports, Roster, walk-ins, the gate
list, unlinking, blocking and moving (§9 r12, r32-36, §10, §17.10-12), and
the console's live-update stream."""

import io
import json
from datetime import datetime, timezone

import openpyxl
import pytest

import config
import db
from conftest import EXPORT, sign
from services import admin, bookings
from test_claims_concurrent import (  # noqa: F401 - fixtures are used by name
    Console, err, roster, secrets_and_limits,
)
from services.roster import COLUMNS  # noqa: E402 - after the fixture named roster

BEFORE = datetime(2026, 9, 17, 4, 0, tzinfo=timezone.utc)


@pytest.fixture()
def world(roster):
    bookings.generate_slots(roster)
    return roster


def pid(conn, handle):
    return conn.execute("SELECT id FROM attendees WHERE handle=?", (handle,)).fetchone()[0]


def data(r):
    body = r.get_json()
    assert body["ok"], body
    return body["data"]


# ---------------------------------------------------------------------------
# Overview
# ---------------------------------------------------------------------------

def test_overview_counts_the_real_list(world):
    world.execute("UPDATE attendees SET payment_status='verified' WHERE handle='heidily'")
    d = data(Console("admin").get("/admin/api/overview"))
    labels = [s["label"] for s in d["stats"]]
    assert labels[:4] == ["Signed up", "Opened the app", "At the event", "Payments verified"]
    assert {"Pastry", "Photo Strip", "Vinyl Crafting", "Escape seats", "Jam slots"} <= set(labels)
    assert d["stats"][0]["value"] == "17"                    # the owner row is a test account
    assert len(d["bars"]) == 21 and d["seats_total"] == 252
    assert [q["what"] for q in d["queue"]][0] == "Refused at the gate"


def stat(d, label):
    return next(s for s in d["stats"] if s["label"] == label)


def test_opening_the_app_and_being_at_the_event_are_counted_apart(world):
    """23 Sep (STATE.md 135). Linking a Telegram account days early is not
    being in the room, and a rehearsal check-in is not either: "at the event"
    starts at the doors on the 24th."""
    overview = lambda: data(Console("admin").get("/admin/api/overview"))  # noqa: E731
    assert stat(overview(), "Opened the app")["value"] == "0"
    assert stat(overview(), "At the event")["value"] == "0"

    # Two people open the app a week early. Nobody is at the event yet.
    world.execute("UPDATE attendees SET tg_user_id=5001 WHERE handle='heidily'")
    world.execute("UPDATE attendees SET tg_user_id=5002 WHERE handle='bananabelles'")
    d = overview()
    assert stat(d, "Opened the app")["value"] == "2"
    assert stat(d, "At the event")["value"] == "0"

    # A rehearsal check-in the day before does not count as being there.
    world.execute("UPDATE attendees SET checked_in_at=? WHERE handle='heidily'",
                  ("2026-09-23T10:00:00+00:00",))
    d = overview()
    assert stat(d, "At the event")["value"] == "0"
    assert "1 before that, not counted" in stat(d, "At the event")["sub"]

    # 3 PM Singapore on the 24th is 07:00 UTC. A minute before still doesn't
    # count; the moment the doors open, it does.
    world.execute("UPDATE attendees SET checked_in_at=? WHERE handle='bananabelles'",
                  ("2026-09-24T06:59:00+00:00",))
    assert stat(overview(), "At the event")["value"] == "0"
    world.execute("UPDATE attendees SET checked_in_at=? WHERE handle='bananabelles'",
                  ("2026-09-24T07:00:00+00:00",))
    d = overview()
    assert stat(d, "At the event")["value"] == "1"
    assert "from 3:00 PM" in stat(d, "At the event")["sub"]
    # Opening the app is unchanged by any of it.
    assert stat(d, "Opened the app")["value"] == "2"


def test_at_the_event_follows_the_doors_setting(world):
    """Move the doors and the counter moves with them — it is not a
    hard-coded 3 PM."""
    world.execute("UPDATE attendees SET checked_in_at=? WHERE handle='heidily'",
                  ("2026-09-24T09:30:00+00:00",))          # 5:30 PM local
    admin.save_settings(world, {"doors_open": "18:00"}, "Max", BEFORE)
    assert stat(data(Console("admin").get("/admin/api/overview")), "At the event")["value"] == "0"
    admin.save_settings(world, {"doors_open": "15:00"}, "Max", BEFORE)
    assert stat(data(Console("admin").get("/admin/api/overview")), "At the event")["value"] == "1"


def test_staff_cannot_see_the_overview(world):
    assert err(Console().get("/admin/api/overview"))["code"] == "FORBIDDEN"


# ---------------------------------------------------------------------------
# Settings
# ---------------------------------------------------------------------------

def test_settings_hide_internal_keys_and_flag_placeholders(world):
    d = data(Console("admin").get("/admin/api/settings"))
    assert "actor_chat_id" not in d["values"]
    assert d["values"]["always_allow"] == "maxi_muslim" and d["readonly"] == ["always_allow"]
    assert set(d["placeholders"]) == {"gm_handle", "actor_handle", "finder_name"}


def test_saving_a_setting_is_logged_with_before_and_after(world):
    a = Console("admin")
    assert data(a.client.put("/admin/api/settings", json={"meeting_point": "Level 5 lobby", "pastry_stock": 40},
                             headers={"X-CSRF-Token": a.csrf}))["saved"] == 2
    assert db.get_setting(world, "meeting_point") == "Level 5 lobby"
    row = world.execute("SELECT * FROM audit_log WHERE action='Setting changed' AND entity_id='pastry_stock'").fetchone()
    assert json.loads(row["details"]) == {"before": 0, "after": 40}
    # A value somebody set is never overwritten by the code's defaults.
    db.seed_settings(world)
    assert db.get_setting(world, "meeting_point") == "Level 5 lobby"


@pytest.mark.parametrize("change, message", [
    ({"nope": 1}, "can't be changed"),
    ({"actor_chat_id": 1}, "can't be changed"),
    ({"doors_open": "5pm"}, "time like"),
    # The hint times went with the script on 23 Sep, so the key itself is
    # no longer a setting anybody can write.
    ({"hint_1": "four"}, "can't be changed"),
    ({"in_app_phone": False}, "can't be changed"),
    ({"capacity": 0}, "at least 1"),
    ({"notify_mode": "loud"}, "one of"),
    ({"test_clock": "yes"}, "on or off"),
    ({"doors_open": "23:00"}, "open before"),
])
def test_bad_settings_are_refused(world, change, message):
    a = Console("admin")
    r = a.client.put("/admin/api/settings", json=change, headers={"X-CSRF-Token": a.csrf})
    assert err(r)["code"] == "VALIDATION_FAILED" and message in err(r)["message"]


def test_event_day_mode_is_gone(world):
    """One Real time / Test time switch replaced it (STATE.md 129)."""
    assert "event_day_mode" not in data(Console("admin").get("/admin/api/settings"))["values"]
    a = Console("admin")
    r = a.client.put("/admin/api/settings", json={"event_day_mode": True}, headers={"X-CSRF-Token": a.csrf})
    assert "can't be changed" in err(r)["message"]


@pytest.mark.parametrize("key", ["doors_message", "last_call"])
def test_the_message_switches_only_take_on_or_off(world, key):
    a = Console("admin")
    r = a.client.put("/admin/api/settings", json={key: "yes"}, headers={"X-CSRF-Token": a.csrf})
    assert "on or off" in err(r)["message"]


def test_test_time_is_the_event_day_at_the_chosen_minute(world, client):
    """It used to keep today's date, which never reached a 24 Sep game."""
    admin.save_settings(world, {"test_clock": True, "test_clock_at": "19:42"}, "Max", BEFORE)
    body = client.get("/healthz").get_json()
    assert body["server_time"].startswith(config.EVENT_DATE + "T19:42")
    session = data(Console("admin").get("/admin/api/session"))
    assert session["test_clock"] is True and session["test_clock_at"] == "19:42"
    admin.save_settings(world, {"test_clock": False}, "Max", BEFORE)
    assert not client.get("/healthz").get_json()["server_time"].startswith(config.EVENT_DATE + "T19:42")


def test_changing_the_timetable_rebuilds_the_games(world):
    out = admin.save_settings(world, {"last_game": "21:25", "capacity": 10}, "Max", BEFORE)
    assert out["schedule"]["escape"]["total"] == 20
    caps = {r[0] for r in world.execute("SELECT capacity FROM slots WHERE room='escape'")}
    assert caps == {10}


def test_handles_in_settings_are_normalised(world):
    admin.save_settings(world, {"gm_handle": "@Ryan_GM "}, "Max", BEFORE)
    assert db.get_setting(world, "gm_handle") == "ryan_gm"


# ---------------------------------------------------------------------------
# Audit, backups, exports
# ---------------------------------------------------------------------------

def test_the_audit_log_filters_and_names_people(world):
    admin.save_settings(world, {"meeting_point": "Lobby"}, "Max", BEFORE)
    world.execute("UPDATE attendees SET payment_status='verified' WHERE handle='heidily'")
    from services import claims
    claims.claim(world, item="photo", actor={"role": "staff", "name": "Wei", "station": "Booth 1"},
                 attendee_id=pid(world, "heidily"))
    a = Console("admin")
    d = data(a.get("/admin/api/audit?action=Photo%20Strip%20handed%20over"))
    assert [e["entity"] for e in d["entries"]] == ["@heidily"]
    assert d["entries"][0]["actor"] == "Wei" and d["entries"][0]["station"] == "Booth 1"
    assert "item: photo" in d["entries"][0]["details"]
    assert data(a.get("/admin/api/audit?search=heidily"))["shown"] >= 1
    setting = data(a.get("/admin/api/audit?actor=Max"))["entries"][0]
    assert setting["details"] == "the front desk → Lobby"
    assert "Wei" in d["actors"] and d["backup"]["count"] == 0


def test_backup_now_makes_a_copy_and_the_timer_knows(world):
    assert admin.backup_due(world)
    d = data(Console("admin").post("/admin/api/backup"))
    assert d["count"] == 1 and d["file"].startswith("app-")
    assert not admin.backup_due(world)
    assert world.execute("SELECT COUNT(*) FROM audit_log WHERE action='Backup made'").fetchone()[0] == 1


def test_the_registrations_export_keeps_paperforms_columns(world):
    h = pid(world, "heidily")
    world.execute("UPDATE attendees SET payment_status='verified' WHERE id=?", (h,))
    from services import claims
    wei = {"role": "staff", "name": "Wei", "station": "Booth 1"}
    claims.check_in(world, actor=wei, attendee_id=h)
    claims.claim(world, item="pastry", actor=wei, attendee_id=h, variant="brownie")
    r = Console("admin").get("/admin/api/export/registrations")
    assert r.status_code == 200 and "attachment" in r.headers["Content-Disposition"]
    ws = openpyxl.load_workbook(io.BytesIO(r.data))["Registrations"]
    rows = list(ws.iter_rows(values_only=True))
    assert list(rows[0]) == COLUMNS
    assert len(rows) == 18                                   # header + 17, no test account
    heidi = next(x for x in rows if x[3] and "heidily" in x[3].lower())
    assert heidi[6] == "Yes" and heidi[8] == "Pastry (Mini Brownie)" and heidi[9]


@pytest.mark.parametrize("kind", ["bookings", "claims", "fallback"])
def test_the_other_exports_download(world, kind):
    r = Console("admin").get(f"/admin/api/export/{kind}")
    assert r.status_code == 200 and len(r.data) > 100


def test_exports_are_admin_only_and_named(world):
    assert err(Console().get("/admin/api/export/claims"))["code"] == "FORBIDDEN"
    assert err(Console("admin").get("/admin/api/export/secrets"))["code"] == "NOT_FOUND"


# ---------------------------------------------------------------------------
# Roster and walk-ins (§10)
# ---------------------------------------------------------------------------

def upload(console, content, name):
    return console.client.post("/admin/api/roster/preview", headers={"X-CSRF-Token": console.csrf},
                               data={"file": (io.BytesIO(content), name)},
                               content_type="multipart/form-data")


def test_re_uploading_the_same_export_changes_nothing(world, data_dir):
    a = Console("admin")
    d = data(upload(a, EXPORT.read_bytes(), EXPORT.name))
    assert d["counts"]["new"] == 0 and d["counts"]["changed"] == 0 and d["counts"]["unchanged"] == 17
    assert data(a.post(f"/admin/api/roster/commit/{d['run_id']}"))["committed"] == 0
    assert err(a.post(f"/admin/api/roster/commit/{d['run_id']}"))["code"] == "VALIDATION_FAILED"
    summary = data(a.get("/admin/api/roster"))
    assert summary["active"] == 17 and len(summary["runs"]) == 2


def test_a_file_that_is_not_an_export_is_refused(world, data_dir):
    a = Console("admin")
    assert "xlsx" in err(upload(a, b"hello", "notes.txt"))["message"]
    assert err(upload(a, b"not really a spreadsheet", "fake.xlsx"))["code"] == "VALIDATION_FAILED"


def test_a_walk_in_gets_a_pass_and_can_open_the_app(world, client):
    a = Console("admin")
    d = data(a.post("/admin/api/people", {"name": " Sam  Lee ", "handle": "@Sam_Walkin", "email": ""}))
    assert d["name"] == "Sam Lee" and d["handle"] == "sam_walkin" and d["pass_code"]
    r = client.post("/api/session", headers={"Authorization": "tma " + sign(7001, "sam_walkin")}, json={})
    assert r.get_json()["ok"]
    assert err(a.post("/admin/api/people", {"name": "Sam", "handle": "sam_walkin"}))["code"] == "HANDLE_TAKEN"
    assert err(a.post("/admin/api/people", {"name": "Sam", "handle": "a b"}))["code"] == "VALIDATION_FAILED"


# ---------------------------------------------------------------------------
# The gate list and unlinking (§9 rules 3 and 7)
# ---------------------------------------------------------------------------

def test_linking_a_refused_account_lets_them_in(world, client):
    auth = {"Authorization": "tma " + sign(8001, "heidi_new")}
    assert err(client.post("/api/session", headers=auth, json={}))["code"] == "NOT_ON_LIST"
    client.post("/api/session/claimed-handle", headers=auth, json={"claimed_handle": "heidily"})
    a = Console("admin")
    attempts = data(a.get("/admin/api/gate-attempts"))["attempts"]
    assert {x["tg_username"] for x in attempts} == {"heidi_new"}
    heidi = pid(world, "heidily")
    assert data(a.post(f"/admin/api/gate-attempts/{attempts[0]['id']}/resolve",
                       {"action": "link", "attendee_id": heidi}))["resolved"] == "link"
    assert client.post("/api/session", headers=auth, json={}).get_json()["ok"]
    assert data(a.get("/admin/api/gate-attempts"))["attempts"] == []      # both entries cleared


def test_a_link_never_takes_over_someone_elses_account(world, client):
    world.execute("UPDATE attendees SET tg_user_id=1 WHERE handle='heidily'")
    client.post("/api/session", headers={"Authorization": "tma " + sign(8002, "stranger1")}, json={})
    a = Console("admin")
    attempt = data(a.get("/admin/api/gate-attempts"))["attempts"][0]
    r = a.post(f"/admin/api/gate-attempts/{attempt['id']}/resolve",
               {"action": "link", "attendee_id": pid(world, "heidily")})
    assert err(r)["code"] == "LINKED_ELSEWHERE"
    assert data(a.post(f"/admin/api/gate-attempts/{attempt['id']}/resolve", {"action": "dismiss"}))


def test_unlink_needs_a_reason(world):
    world.execute("UPDATE attendees SET tg_user_id=55 WHERE handle='heidily'")
    a = Console("admin")
    path = f"/admin/api/people/{pid(world, 'heidily')}/unlink"
    assert err(a.post(path, {"reason": " "}))["code"] == "VALIDATION_FAILED"
    assert data(a.post(path, {"reason": "wrong account"}))["unlinked"]
    assert world.execute("SELECT tg_user_id FROM attendees WHERE handle='heidily'").fetchone()[0] is None


# ---------------------------------------------------------------------------
# Blocking and moving
# ---------------------------------------------------------------------------

def first_games(conn, n=2):
    return conn.execute("SELECT * FROM slots WHERE room='escape' ORDER BY starts_at LIMIT ?", (n,)).fetchall()


def test_blocking_a_game_hides_it_and_names_who_is_still_booked(world):
    g1, _ = first_games(world)
    bookings.book(world, pid(world, "heidily"), g1["id"], [], BEFORE)
    a = Console("admin")
    path = f"/admin/api/slots/{g1['id']}/block"
    assert err(a.post(path, {"blocked": True}))["code"] == "VALIDATION_FAILED"      # reason needed
    d = data(a.post(path, {"blocked": True, "reason": "actor ill"}))
    assert d["still_booked"] == ["@heidily"]
    board = bookings.board(world, pid(world, "bananabelles"), BEFORE)
    assert board["slots"][0]["status"] == "blocked"
    assert data(a.post(path, {"blocked": False}))["blocked"] is False


def test_moving_someone_changes_their_game_and_tells_them(world):
    g1, g2 = first_games(world)
    h = pid(world, "heidily")
    bookings.book(world, h, g1["id"], [], BEFORE)
    a = Console("admin")
    d = data(a.post("/admin/api/escape/move", {"attendee_id": h, "slot_id": g2["id"], "reason": "late"}))
    assert d["moved_to"].startswith("2026-09-24T15:25")
    assert bookings.active_booking(world, h)["slot_id"] == g2["id"]
    assert world.execute("SELECT COUNT(*) FROM notifications WHERE kind='moved'").fetchone()[0] == 1
    assert err(a.post("/admin/api/escape/move", {"attendee_id": h, "slot_id": g2["id"]}))["code"] \
        == "VALIDATION_FAILED"


def test_moving_into_a_full_game_is_refused(world):
    g1, g2 = first_games(world)
    world.execute("UPDATE slots SET capacity=0 WHERE id=?", (g2["id"],))
    h = pid(world, "heidily")
    bookings.book(world, h, g1["id"], [], BEFORE)
    r = Console("admin").post("/admin/api/escape/move", {"attendee_id": h, "slot_id": g2["id"]})
    assert err(r)["code"] == "SLOT_FULL"


# ---------------------------------------------------------------------------
# Live updates
# ---------------------------------------------------------------------------

def test_the_live_stream_says_when_something_changed(world, monkeypatch):
    import app as webapp
    monkeypatch.setattr(webapp, "LIVE_TICK_SECONDS", 0.01)
    monkeypatch.setattr(config, "LIVE_STREAM_SECONDS", 0.05)
    r = Console("gm", "Ryan", "GM").get("/admin/api/live")
    assert r.status_code == 200 and r.mimetype == "text/event-stream"
    text = r.get_data(as_text=True)
    assert text.startswith("retry: 3000") and "data: {\"v\": [" in text


def test_the_live_stream_needs_a_sign_in(world, client):
    assert err(client.get("/admin/api/live"))["code"] == "SIGNED_OUT"
